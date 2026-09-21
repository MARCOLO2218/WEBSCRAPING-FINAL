from openpyxl import Workbook
from catalog_api.catalog import excel_rows, merge, compare


def test_regional_excel_never_mixes_countries_or_currencies(tmp_path):
    path = tmp_path / "regional.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Precios FACENCO"
    for i, header in enumerate(["pais", "moneda", "producto", "precio_regular", "precio_oferta"], 1):
        sheet.cell(4, i, header)
    pairs = [("GT", "GTQ"), ("HN", "HNL"), ("SV", "USD"), ("NC", "NIO"),
             (None, "GTQ"), ("GT", "HNL"), ("GT", None), ("NI", "NIO")]
    for r, (country, currency) in enumerate(pairs, 5):
        for c, value in enumerate([country, currency, "Cama", r * 100, r * 90], 1):
            sheet.cell(r, c, value)
    book.save(path)
    loaded = excel_rows([], path)
    assert len(loaded) == 1
    assert loaded[0]["precio_regular"] == "Q500"


def test_excel_price_merge_keeps_identity_and_reference(tmp_path):
    path = tmp_path / "prices.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Precios FACENCO"
    for column, value in enumerate(["producto", "precio_regular", "precio_oferta", "activo", "codigo_producto"], 1):
        sheet.cell(4, column, value)
    for column, value in enumerate(["Colchón", 3000, 2500, "SI", "C1"], 1):
        sheet.cell(5, column, value)
    sheet.cell(6, 1, "Inactivo")
    sheet.cell(6, 4, "NO")
    book.save(path)
    source = [{"id": "11", "run_id": "20", "producto": "COLCHON", "sitio_fuente": "FACENCO"}]
    loaded = excel_rows(source, path)
    assert len(loaded) == 1
    assert loaded[0]["precio_oferta"] == "Q2,500"
    result = compare(merge(source, loaded))
    assert len(result) == 1
    assert result[0]["id"] == "11"
    assert result[0]["precio_numero"] == 2500
    assert result[0]["etiqueta_diferencia"] == "Igual a FACENCO"
